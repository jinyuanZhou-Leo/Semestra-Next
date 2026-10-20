"""One reader per file type; each returns a list of pages of markdown (M07 §3.1)."""

import csv
import io
import re
import zipfile
from collections.abc import Iterable, Sequence
from dataclasses import dataclass, field
from html.parser import HTMLParser
from typing import Any, ClassVar, cast

import pymupdf
import pymupdf4llm

from semestra.modules.files.parse.detect import Detected, decode_text

MAX_PAGES = 500
OCR_MIN_CHARS = 50
LOGICAL_PAGE_CHARS = 3000
TABLE_ROWS_PER_PAGE = 100
MAX_TABLE_ROWS = 20_000
# Office files are zip archives; the readers inflate whole XML parts, so a small file that
# unpacks to hundreds of megabytes is refused before it is opened.
MAX_OFFICE_UNPACKED_BYTES = 100_000_000
PARSER_VERSION = "1"


class ParseError(Exception):
    """The file cannot be read; `code` is the error code stored on the item."""

    def __init__(self, code: str, **params: object) -> None:
        super().__init__(code)
        self.code = code
        self.params = params


@dataclass
class ParsedPage:
    markdown: str
    needs_ocr: bool = False


@dataclass
class ParsedDocument:
    pages: list[ParsedPage] = field(default_factory=lambda: [])

    @property
    def char_count(self) -> int:
        return sum(len(p.markdown) for p in self.pages)


def _split_long(block: str, size: int) -> list[str]:
    """Cuts a paragraph longer than `size` at line ends, and lines longer than `size` at
    whitespace (or anywhere), so no page grows past `size` (text without blank lines)."""
    if len(block) <= size:
        return [block]
    parts: list[str] = []
    current = ""
    for line in block.split("\n"):
        while len(line) > size:
            cut = line.rfind(" ", size // 2, size)
            cut = size if cut <= 0 else cut
            if current:
                parts.append(current)
                current = ""
            parts.append(line[:cut])
            line = line[cut:].lstrip(" ")
        if current and len(current) + 1 + len(line) > size:
            parts.append(current)
            current = line
        else:
            current = f"{current}\n{line}" if current else line
    if current:
        parts.append(current)
    return parts


def paginate(blocks: Sequence[str], size: int = LOGICAL_PAGE_CHARS) -> list[str]:
    """Groups paragraphs into logical pages of about `size` characters; longer paragraphs are
    cut so a page never exceeds it."""
    pages: list[str] = []
    current: list[str] = []
    length = 0
    for block in (part for b in blocks for part in _split_long(b, size)):
        if current and length + len(block) > size:
            pages.append("\n\n".join(current))
            current, length = [], 0
        current.append(block)
        length += len(block) + 2
    if current:
        pages.append("\n\n".join(current))
    return pages


def _paragraphs(text: str) -> list[str]:
    return [p.strip() for p in re.split(r"\n\s*\n", text.replace("\r\n", "\n")) if p.strip()]


def _cell(value: object) -> str:
    text = "" if value is None else str(value)
    return text.replace("|", "\\|").replace("\r", "").replace("\n", " ").strip()


def table_markdown(rows: Sequence[Sequence[object]]) -> str:
    cells = [[_cell(c) for c in row] for row in rows]
    if not cells:
        return ""
    width = max(len(r) for r in cells)
    cells = [r + [""] * (width - len(r)) for r in cells]
    lines = ["| " + " | ".join(cells[0]) + " |", "|" + " --- |" * width]
    lines += ["| " + " | ".join(r) + " |" for r in cells[1:]]
    return "\n".join(lines)


def _table_pages(title: str | None, rows: list[list[object]]) -> list[ParsedPage]:
    rows = [r for r in rows if any(str(c).strip() for c in r if c is not None)]
    if not rows:
        return [ParsedPage(f"## {title}" if title else "")] if title else []
    header, body = rows[0], rows[1:]
    chunks = [body[i : i + TABLE_ROWS_PER_PAGE] for i in range(0, len(body), TABLE_ROWS_PER_PAGE)]
    pages: list[ParsedPage] = []
    for chunk in chunks or [[]]:
        text = table_markdown([header, *chunk])
        pages.append(ParsedPage(f"## {title}\n\n{text}" if title else text))
    return pages


def _text_pages(text: str) -> list[ParsedPage]:
    return [ParsedPage(p) for p in paginate(_paragraphs(text))]


# --- PDF ------------------------------------------------------------------------------------


def _pdf(data: bytes, max_pages: int) -> ParsedDocument:
    try:
        doc = pymupdf.open(stream=data, filetype="pdf")
    except Exception as exc:
        raise ParseError("files.parse_failed") from exc
    with doc:
        if doc.needs_pass:
            raise ParseError("files.encrypted_pdf")
        if doc.page_count > max_pages:
            raise ParseError("files.too_many_pages", max=max_pages)
        try:
            headers = pymupdf4llm.IdentifyHeaders(doc)  # pyright: ignore[reportArgumentType]
            pages: list[ParsedPage] = []
            for index in range(doc.page_count):
                result = pymupdf4llm.to_markdown(
                    doc, pages=[index], hdr_info=headers, show_progress=False
                )
                markdown = result.strip() if isinstance(result, str) else ""
                sparse = len(re.sub(r"\s+", "", markdown)) < OCR_MIN_CHARS
                pages.append(ParsedPage(markdown, sparse and bool(doc[index].get_images())))
        except ParseError:
            raise
        except Exception as exc:
            raise ParseError("files.parse_failed") from exc
    return ParsedDocument(pages)


# --- Office ---------------------------------------------------------------------------------


def _docx(data: bytes) -> ParsedDocument:
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    document = docx.Document(io.BytesIO(data))
    pages: list[list[str]] = [[]]
    length = 0

    def new_page() -> None:
        nonlocal length
        if pages[-1]:
            pages.append([])
        length = 0

    for element in document.element.body.iterchildren():
        tag = element.tag.rsplit("}", 1)[-1]
        if tag == "tbl":
            rows = [[c.text for c in row.cells] for row in Table(element, document).rows]  # pyright: ignore[reportArgumentType]
            block = table_markdown(rows)
        elif tag == "p":
            paragraph = Paragraph(element, document)  # pyright: ignore[reportArgumentType]
            if element.xpath(".//w:lastRenderedPageBreak"):  # pyright: ignore[reportAttributeAccessIssue]
                new_page()
            text = paragraph.text.strip()
            style = (paragraph.style.name if paragraph.style is not None else "") or ""
            block = ""
            if text:
                heading = re.fullmatch(r"Heading (\d)", style)
                if heading:
                    block = f"{'#' * min(int(heading[1]), 6)} {text}"
                elif style == "Title":
                    block = f"# {text}"
                elif style.startswith("List"):
                    block = f"- {text}"
                else:
                    block = text
        else:
            continue
        if block:
            if length and length + len(block) > LOGICAL_PAGE_CHARS:
                new_page()
            pages[-1].append(block)
            length += len(block)
        if tag == "p" and element.xpath('.//w:br[@w:type="page"]'):  # pyright: ignore[reportAttributeAccessIssue]
            new_page()
    return ParsedDocument([ParsedPage("\n\n".join(p)) for p in pages if p])


def _pptx(data: bytes) -> ParsedDocument:
    from pptx import Presentation

    pages: list[ParsedPage] = []
    for slide in Presentation(io.BytesIO(data)).slides:
        title_shape = slide.shapes.title
        title = title_shape.text_frame.text.strip() if title_shape is not None else ""
        blocks = [f"# {title}"] if title else []
        for shape in cast("Iterable[Any]", slide.shapes):
            if shape is title_shape:
                continue
            if shape.has_text_frame:
                lines = [
                    f"{'  ' * paragraph.level}- {paragraph.text.strip()}"
                    for paragraph in shape.text_frame.paragraphs
                    if paragraph.text.strip()
                ]
                if lines:
                    blocks.append("\n".join(lines))
            elif shape.has_table:
                blocks.append(
                    table_markdown([[c.text for c in row.cells] for row in shape.table.rows])
                )
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame
            if notes is not None and notes.text.strip():
                blocks.append(f"Notes: {notes.text.strip()}")
        pages.append(ParsedPage("\n\n".join(blocks)))
    return ParsedDocument(pages)


def _xlsx(data: bytes) -> ParsedDocument:
    import openpyxl

    workbook = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    pages: list[ParsedPage] = []
    try:
        for sheet in workbook.worksheets:
            rows: list[list[object]] = []
            for row in sheet.iter_rows(values_only=True):
                rows.append(list(row))
                if len(rows) >= MAX_TABLE_ROWS:
                    break
            pages += _table_pages(sheet.title, rows)
    finally:
        workbook.close()
    return ParsedDocument(pages)


def _csv(data: bytes, tab: bool) -> ParsedDocument:
    text = decode_text(data) or ""
    sample = text[:4096]
    if tab:
        delimiter = "\t"
    else:
        try:
            delimiter = csv.Sniffer().sniff(sample, delimiters=",;\t").delimiter
        except csv.Error:
            delimiter = ","
    rows: list[list[object]] = []
    for row in csv.reader(io.StringIO(text), delimiter=delimiter):
        rows.append(list(row))
        if len(rows) >= MAX_TABLE_ROWS:
            break
    return ParsedDocument(_table_pages(None, rows))


# --- text, markdown, HTML, images -----------------------------------------------------------


class _HtmlText(HTMLParser):
    SKIP: ClassVar[set[str]] = {"script", "style", "noscript", "template", "head", "svg"}
    BLOCK: ClassVar[set[str]] = {
        "p",
        "div",
        "br",
        "tr",
        "section",
        "article",
        "ul",
        "ol",
        "table",
        "blockquote",
    }

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.parts: list[str] = []
        self._skip = 0

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        del attrs
        if tag in self.SKIP:
            self._skip += 1
        elif re.fullmatch(r"h[1-6]", tag):
            self.parts.append(f"\n\n{'#' * int(tag[1])} ")
        elif tag == "li":
            self.parts.append("\n- ")
        elif tag in self.BLOCK:
            self.parts.append("\n\n")
        elif tag in ("td", "th"):
            self.parts.append(" | ")

    def handle_endtag(self, tag: str) -> None:
        if tag in self.SKIP:
            self._skip = max(0, self._skip - 1)
        elif re.fullmatch(r"h[1-6]", tag):
            self.parts.append("\n\n")

    def handle_data(self, data: str) -> None:
        if not self._skip:
            self.parts.append(data)


def _html(data: bytes) -> ParsedDocument:
    parser = _HtmlText()
    parser.feed(decode_text(data) or "")
    text = re.sub(r"[ \t]+", " ", "".join(parser.parts))
    return ParsedDocument(_text_pages(re.sub(r"\n[ \t]*\n(?:[ \t]*\n)+", "\n\n", text)))


def parse(data: bytes, detected: Detected, *, max_pages: int = MAX_PAGES) -> ParsedDocument:
    """Reads the file. Raises `ParseError` for files that cannot be read."""
    try:
        result = _dispatch(data, detected, max_pages)
    except ParseError:
        raise
    except Exception as exc:
        raise ParseError("files.parse_failed") from exc
    if len(result.pages) > max_pages:
        raise ParseError("files.too_many_pages", max=max_pages)
    return result


def _check_unpacked_size(data: bytes) -> None:
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        if sum(info.file_size for info in archive.infolist()) > MAX_OFFICE_UNPACKED_BYTES:
            raise ParseError("files.too_complex")


def _dispatch(data: bytes, detected: Detected, max_pages: int) -> ParsedDocument:
    if detected.kind in ("docx", "pptx", "xlsx"):
        _check_unpacked_size(data)
    match detected.kind:
        case "pdf":
            return _pdf(data, max_pages)
        case "docx":
            return _docx(data)
        case "pptx":
            return _pptx(data)
        case "xlsx":
            return _xlsx(data)
        case "csv":
            return _csv(data, tab=False)
        case "html":
            return _html(data)
        case "text" | "markdown":
            return ParsedDocument(_text_pages(decode_text(data) or ""))
        case "image":
            return ParsedDocument([ParsedPage("", needs_ocr=True)])
        case _:
            raise ParseError("files.unsupported_type")


def guess_language(text: str) -> str | None:
    """ "zh" when the text is mostly CJK, "en" when it is mostly Latin letters."""
    sample = text[:20_000]
    cjk = len(re.findall(r"[\u3400-\u9fff]", sample))
    latin = len(re.findall(r"[A-Za-z]", sample))
    if cjk < 6 and latin < 20:
        return None
    return "zh" if cjk > latin * 0.3 else "en"
