"""Type detection and the per-format readers (M07 §3.1, §9 snapshot-style checks)."""

import io
import os
import signal
import sys
from contextlib import aclosing
from typing import Any

import pytest

from semestra.modules.files.parse import ParseError, detect, parse
from semestra.modules.files.parse.readers import LOGICAL_PAGE_CHARS, guess_language
from semestra.modules.files.tests.conftest import SYLLABUS, make_docx, make_pdf, make_png


def test_detection_uses_the_content_not_the_name():
    assert detect(make_pdf(["x"]), "notes.txt").kind == "pdf"
    assert detect(make_png(), "photo.pdf").mime == "image/png"
    assert detect(make_docx([("Normal", "hi")]), "a.bin").kind == "docx"
    assert detect(b"BEGIN:VCALENDAR\r\nEND:VCALENDAR", "x.ics").kind == "ics"
    assert detect(b"<!doctype html><p>hi</p>", "page").kind == "html"
    assert detect(b"a,b\n1,2", "data.csv").kind == "csv"
    assert detect(b"# Title", "readme.md").mime == "text/markdown"
    assert detect("纯文本".encode("gb18030"), "a.txt").kind == "text"
    assert detect(b"\x00\x01\x02binary", "a.txt").kind == "other"
    assert detect(b"<svg xmlns='http://www.w3.org/2000/svg'><script/></svg>", "x.svg").mime == (
        "text/plain"
    )


def test_digital_pdf_gives_one_markdown_page_per_page():
    doc = parse(make_pdf(SYLLABUS), detect(make_pdf(SYLLABUS)))
    assert len(doc.pages) == 3
    assert "APS105 Computer Fundamentals" in doc.pages[0].markdown
    assert "30 percent" in doc.pages[1].markdown
    assert "期中考试时间" in doc.pages[2].markdown
    assert not any(p.needs_ocr for p in doc.pages)


def test_scanned_pdf_pages_are_marked_for_ocr():
    data = make_pdf(
        ["typed page with enough words to count as a digital text page", "scan"], scanned={1}
    )
    doc = parse(data, detect(data))
    assert [p.needs_ocr for p in doc.pages] == [False, True]


def test_encrypted_and_oversized_pdfs_are_refused():
    import pymupdf

    source = pymupdf.open(stream=make_pdf(["secret"]), filetype="pdf")
    locked = source.tobytes(
        encryption=pymupdf.PDF_ENCRYPT_AES_256,  # pyright: ignore[reportAttributeAccessIssue]
        user_pw="pw",
        owner_pw="pw",
    )
    with pytest.raises(ParseError) as exc:
        parse(locked, detect(locked))
    assert exc.value.code == "files.encrypted_pdf"

    many = make_pdf(["p"] * 4)
    with pytest.raises(ParseError) as exc:
        parse(many, detect(many), max_pages=3)
    assert exc.value.code == "files.too_many_pages"


def test_docx_headings_and_paragraphs():
    data = make_docx(
        [
            ("Heading 1", "Course outline"),
            ("Normal", "Weekly quizzes."),
            ("Heading 2", "Late policy"),
            ("Normal", "10% per day."),
        ]
    )
    doc = parse(data, detect(data))
    text = doc.pages[0].markdown
    assert "# Course outline" in text
    assert "## Late policy" in text
    assert "10% per day." in text


def test_long_docx_is_split_into_logical_pages():
    data = make_docx([("Normal", "word " * 200) for _ in range(20)])
    assert len(parse(data, detect(data)).pages) >= 3


def test_text_without_blank_lines_is_still_split_into_pages():
    lines = "\n".join(f"DTSTART:2026091{n % 10}T100000 SUMMARY:Lecture {n}" for n in range(400))
    one_line = "word " * 2000
    for data in (lines.encode(), one_line.encode()):
        doc = parse(data, detect(data, "notes.txt"))
        assert len(doc.pages) > 1
        assert max(len(p.markdown) for p in doc.pages) <= LOGICAL_PAGE_CHARS
        assert " ".join(p.markdown for p in doc.pages).split() == data.decode().split()


def test_pptx_has_one_page_per_slide_with_notes():
    from pptx import Presentation

    deck = Presentation()
    for title, body in [("Week 1", "Intro"), ("Week 2", "Pointers")]:
        slide = deck.slides.add_slide(deck.slide_layouts[1])
        slide.shapes.title.text = title  # pyright: ignore[reportOptionalMemberAccess]
        slide.placeholders[1].text = body  # pyright: ignore[reportAttributeAccessIssue]
        slide.notes_slide.notes_text_frame.text = f"say {body}"  # pyright: ignore[reportOptionalMemberAccess]
    out = io.BytesIO()
    deck.save(out)
    doc = parse(out.getvalue(), detect(out.getvalue()))
    assert len(doc.pages) == 2
    assert doc.pages[1].markdown.startswith("# Week 2")
    assert "- Pointers" in doc.pages[1].markdown
    assert "Notes: say Pointers" in doc.pages[1].markdown


def test_spreadsheets_become_markdown_tables():
    import openpyxl

    book = openpyxl.Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.title = "Marks"
    sheet.append(["Item", "Weight"])
    sheet.append(["Midterm", 30])
    out = io.BytesIO()
    book.save(out)
    doc = parse(out.getvalue(), detect(out.getvalue()))
    assert doc.pages[0].markdown == "## Marks\n\n| Item | Weight |\n| --- | --- |\n| Midterm | 30 |"

    csv_doc = parse(b"a;b\n1;2\n", detect(b"a;b\n1;2\n", "x.csv"))
    assert "| 1 | 2 |" in csv_doc.pages[0].markdown


def test_spreadsheets_drop_empty_columns_and_long_tables_are_cut():
    import openpyxl

    book = openpyxl.Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.append(["Item", "Weight"])
    sheet.cell(row=1, column=60, value=" ")  # formatted far to the right, but empty
    for n in range(450):
        sheet.append([f"Quiz {n}", 1])
    out = io.BytesIO()
    book.save(out)
    doc = parse(out.getvalue(), detect(out.getvalue()), max_pages=3)
    assert len(doc.pages) == 3
    assert doc.pages[0].markdown.splitlines()[2] == "| Item | Weight |"


def test_html_keeps_headings_and_drops_scripts():
    data = (
        b"<html><head><title>t</title><script>alert(1)</script></head>"
        b"<body><h2>Rules</h2><p>Be kind</p><ul><li>One</li></ul></body></html>"
    )
    text = parse(data, detect(data)).pages[0].markdown
    assert "## Rules" in text
    assert "Be kind" in text
    assert "- One" in text
    assert "alert" not in text
    table = b"<p>See <a href='https://q.ca'>Quercus</a></p><table><tr><th>A</th></tr><tr><td>1</td></tr></table>"
    text = parse(table, detect(table, "x.html")).pages[0].markdown
    assert "[Quercus](https://q.ca)" in text
    assert "| A |" in text


def test_text_encodings_are_guessed():
    from semestra.modules.files.parse.detect import decode_text

    notice = "期中考试时间和地点，请准时到场。"  # noqa: RUF001
    assert decode_text(notice.encode("gbk")) == notice
    assert decode_text("第一週課程".encode("big5")) == "第一週課程"
    # A Western CSV is not read as Chinese (what trying gb18030 first did).
    western = decode_text("Müller, Café, résumé".encode("cp1252")) or ""
    assert western.startswith("M")
    assert not any("\u4e00" <= ch <= "\u9fff" for ch in western)


def test_images_are_one_page_waiting_for_ocr():
    data = make_png()
    doc = parse(data, detect(data))
    assert len(doc.pages) == 1
    assert doc.pages[0].needs_ocr


def test_unreadable_files_raise():
    with pytest.raises(ParseError) as exc:
        parse(b"\x00binary", detect(b"\x00binary"))
    assert exc.value.code == "files.unsupported_type"
    broken = b"%PDF-1.7 broken"
    with pytest.raises(ParseError) as exc:
        parse(broken, detect(broken))
    assert exc.value.code == "files.parse_failed"


def test_language_guess():
    assert guess_language("这门课程的期中考试在十月") == "zh"
    assert guess_language("The midterm exam is in October this year") == "en"
    assert guess_language("123") is None


def test_office_files_that_unpack_too_large_are_refused(monkeypatch: pytest.MonkeyPatch):
    from semestra.modules.files.parse import readers

    data = make_docx([("Normal", "fine")])
    monkeypatch.setattr(readers, "MAX_OFFICE_UNPACKED_BYTES", 1000)
    with pytest.raises(ParseError) as exc:
        parse(data, detect(data))
    assert exc.value.code == "files.too_complex"


def test_page_images_are_bounded_whatever_the_page_shape():
    import pymupdf
    from PIL import Image

    from semestra.modules.files.parse.render import (
        OCR_MAX_PIXELS,
        THUMB_MAX_HEIGHT,
        THUMB_WIDTH,
        PageRenderer,
    )

    doc = pymupdf.open()
    doc.new_page(width=40, height=14400)
    doc.new_page(width=14400, height=14400)
    renderer = PageRenderer(doc.tobytes(), pdf=True)
    for index in (0, 1):
        thumb = Image.open(io.BytesIO(renderer.thumbnail(index)))
        assert thumb.width <= THUMB_WIDTH
        assert thumb.height <= THUMB_MAX_HEIGHT
        ocr = Image.open(io.BytesIO(renderer.ocr_image(index)))
        assert ocr.width * ocr.height <= OCR_MAX_PIXELS * 1.01
    renderer.close()

    tall = Image.new("RGB", (100, 50_000), "white")
    out = io.BytesIO()
    tall.save(out, "PNG")
    image = PageRenderer(out.getvalue(), pdf=False)
    thumb = Image.open(io.BytesIO(image.thumbnail(0)))
    assert thumb.height <= THUMB_MAX_HEIGHT


# Stand-ins for the parsing child process; module-level so the child process can import them.


def hang(*_args: object) -> None:
    import time

    time.sleep(60)


def die(*_args: object) -> None:
    import os
    import signal

    os.kill(os.getpid(), signal.SIGKILL)


def run_out_of_memory(conn: object, *_args: object) -> None:
    from semestra.modules.files.parse import isolated

    def explode(*_a: object, **_k: object) -> None:
        raise MemoryError

    isolated.parse = explode  # pyright: ignore[reportAttributeAccessIssue]
    isolated._child(conn, *_args)  # pyright: ignore[reportPrivateUsage, reportArgumentType]


@pytest.mark.parametrize("target", [hang, die, run_out_of_memory])
async def test_parsing_that_hangs_crashes_or_exhausts_memory_is_too_complex(
    target: object, monkeypatch: pytest.MonkeyPatch
):
    from semestra.modules.files.parse import isolated

    monkeypatch.setattr(isolated, "_child", target)
    data = make_pdf(["hello"])
    with pytest.raises(ParseError) as exc:
        async for _ in isolated.parse_isolated(
            data,
            detect(data),
            max_pages=10,
            thumbnail_pages=10,
            time_limit=1,
            memory_limit=2**31,
        ):
            pass
    assert exc.value.code == "files.too_complex"


async def test_parsing_in_a_child_process_yields_the_document_and_previews():
    from semestra.modules.files.parse.isolated import Parsed, Thumbnail, parse_isolated

    data = make_pdf(SYLLABUS)
    messages = [
        m
        async for m in parse_isolated(
            data,
            detect(data),
            max_pages=10,
            thumbnail_pages=2,
            time_limit=60,
            memory_limit=2**31,
        )
    ]
    assert isinstance(messages[0], Parsed)
    assert len(messages[0].document.pages) == 3
    assert [m.index for m in messages[1:] if isinstance(m, Thumbnail)] == [0, 1]


async def test_page_images_render_in_a_child_process():
    from semestra.modules.files.parse.isolated import render_isolated

    data = make_pdf(["one", "two", "three"], scanned={1})
    limits = {"time_limit": 60, "memory_limit": 2**31}
    chosen = await render_isolated(data, pdf=True, pages=[1], **limits)
    first = await render_isolated(data, pdf=True, pages=5, **limits)
    assert len(chosen) == 1
    assert chosen[0].startswith(b"\x89PNG")
    assert len(first) == 3  # at most the pages there are


def _killed(conn: object) -> None:
    del conn
    os.kill(os.getpid(), signal.SIGKILL)  # what the memory limit or a crash in C code does


def _exits(conn: object) -> None:
    del conn
    sys.exit(1)


async def test_a_child_that_dies_or_runs_out_of_time_is_reported():
    from semestra.modules.files.parse import isolated

    async def run(target: Any, time_limit: float = 60) -> None:
        async with aclosing(isolated._run(target, (), time_limit)) as messages:  # pyright: ignore[reportPrivateUsage]
            async for _ in messages:
                pass

    with pytest.raises(ParseError) as killed:
        await run(_killed)
    assert killed.value.code == "files.too_complex"
    with pytest.raises(ParseError) as exited:
        await run(_exits)
    assert exited.value.code == "files.parse_failed"
    data = make_pdf(["one", "two"], scanned={1})
    with pytest.raises(ParseError) as late:  # the deadline holds even while pages arrive
        await isolated.render_isolated(data, pdf=True, pages=2, time_limit=0, memory_limit=2**31)
    assert late.value.code == "files.too_complex"
